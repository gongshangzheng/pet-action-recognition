"""Create a Chinese report, sortable standalone HTML, Excel and charts from a batch."""
from __future__ import annotations
import argparse
import base64
import csv
import html
import json
from pathlib import Path

NAMES={'c2d-resnet50':'C2D-R50','c3d-sports1m':'C3D','csn-ircsn152':'CSN-IR152',
 'i3d-resnet50':'I3D-R50','mvit-small':'MViT-Small','r2plus1d-resnet34':'R(2+1)D-R34',
 'slowfast-resnet50':'SlowFast-R50','slowonly-resnet50':'SlowOnly-R50',
 'swin-tiny':'Swin-Base [ID: swin-tiny]','tanet-resnet50':'TANet-R50','timesformer-divst':'TimeSformer',
 'tin-resnet50':'TIN-R50','tpn-slowonly-r50':'TPN-SlowOnly-R50','trn-resnet50':'TRN-R50',
 'tsm-resnet50':'TSM-R50','tsn-resnet50':'TSN-R50','videomae-base':'VideoMAE-Base',
 'videomaev2-base':'VideoMAEv2-Base'}
ZH={'locomotion':'移动','jump':'跳跃','eating':'进食','drinking':'饮水','grooming':'梳理/舔毛',
    'still_rest':'静止/休息','social_interaction':'社会互动'}

def fmt(value,digits=2):
    return '—' if value is None else f'{value:.{digits}f}'

def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('batch');a=ap.parse_args()
    p=Path(a.batch)
    data=json.loads((p/'comparison.json').read_text(encoding='utf-8'))
    manifest=json.loads((p/'manifest.json').read_text(encoding='utf-8'))
    env=json.loads((p/'environment.json').read_text(encoding='utf-8'))
    good=sorted([r for r in data if r['status']=='completed'],key=lambda r:r['classification']['macro_f1'],reverse=True)
    other=[r for r in data if r['status']!='completed']
    dist=list(csv.DictReader((p/'dataset/class_distribution.csv').open(encoding='utf-8')))
    video_meta={r['relative_path']:r for r in csv.DictReader((p/'dataset/video_manifest.csv').open(encoding='utf-8'))}
    split_hashes={}
    for split in ['train','val','test']:
        paths=[line.rsplit(maxsplit=1)[0] for line in (p/'dataset'/f'{split}_public.txt').read_text().splitlines() if line.strip()]
        split_hashes[split]={video_meta[path]['sha256'] for path in paths}
    duplicate_hashes={a+'_'+b:sorted(split_hashes[a]&split_hashes[b]) for a,b in [('train','val'),('train','test'),('val','test')]}
    (p/'dataset/content_overlap_audit.json').write_text(json.dumps(dict(source='stored video_manifest SHA256, not freshly rehashed',cross_split_duplicates=duplicate_hashes),indent=2),encoding='utf-8')
    all_class=[]
    for r in good:
        for row in r['classification']['per_class']:
            all_class.append(dict(model_id=r['model_id'],model=NAMES.get(r['model_id'],r['model_id']),**row))
    if all_class:
        with (p/'per_class_all_models.csv').open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fieldnames=list(all_class[0]));w.writeheader();w.writerows(all_class)
    # Independently audit sample counts, stored confusion matrices and Top-1.
    audits=[]
    for r in good:
        lines=[json.loads(x) for x in (p/'models'/r['model_id']/'predictions.jsonl').read_text().splitlines()]
        cm=[[0]*7 for _ in range(7)]
        for x in lines: cm[x['target']][x['prediction']]+=1
        accuracy=sum(x['correct'] for x in lines)/len(lines)
        passed=len(lines)==217 and cm==r['confusion_matrix'] and abs(accuracy-r['classification']['top1'])<1e-10
        audits.append(dict(model=r['model_id'],predictions=len(lines),metrics_verified=passed))
        if not passed: raise RuntimeError(f'Metric audit failed for {r["model_id"]}')
    (p/'report_audit.json').write_text(json.dumps(audits,indent=2),encoding='utf-8')
    headers=['模型（按 Macro-F1 排序）','Top-1 %','Macro-F1 %','宏平均召回 %','参数 M','GFLOPs†',
             '前向均值 ms','前向 P95 ms','整体均值 ms','测试数']
    table=[]
    for r in good:
        c=r['classification'];fl=r.get('flops',{})
        g=fmt(fl.get('gflops'))+('*' if fl.get('unsupported_ops') else '')
        table.append([NAMES.get(r['model_id'],r['model_id']),fmt(c['top1']*100),fmt(c['macro_f1']*100),
          fmt(c['macro_recall']*100),fmt(r['parameters']/1e6),g,fmt(r['forward_host']['mean_ms']),
          fmt(r['forward_host']['p95_ms']),fmt(r['e2e']['mean_ms']),str(r['successful_samples'])])
    md=['# AnimalKingdom 哺乳动物子集模型对比', '',
        '本报告为本批次重新执行的评测结果；使用已有微调权重，没有重新训练。', '',
        f'完整完成 {len(good)}/{len(data)} 个候选模型；按 Macro-F1 排序。失败和未满足评测条件的模型均单列。', '',
        '## 数据来源与适用范围', '',
        '- 数据集：`pet_action_mammal_v0`，来源为 AnimalKingdom 的严格哺乳动物子集，并重映射为七类单标签分类。',
        '- 训练/验证/测试：1801 / 216 / 217。保留原划分，不使用测试结果选择权重或调参。',
        '- 原 AnimalKingdom 训练划分仍为训练；原测试划分中的筛选样本再划分为本项目验证/测试。',
        '- 这不是原始 AnimalKingdom 全物种、原标签体系的官方基准，也不是家猫五分类测试。',
        '- 七类为移动、跳跃、进食、饮水、梳理/舔毛、静止/休息、社会互动；`other_unknown` 无样本，不作为第八类。',
        '- 片段未经人工视觉复核，继承源标签映射的局限；清单路径不重叠。另按原始视频清单中的 SHA256 检查跨划分相同内容；这不能排除来自同一个长视频的关联片段。',
        '- 基于原始 SHA256 清单的跨划分相同内容数：'+str({k:len(v) for k,v in duplicate_hashes.items()})+'。',
        f'- 测试清单 SHA256：`{manifest["dataset_audit"]["test_sha256"]}`。', '',
        '| 类别 | 训练 | 验证 | 测试 |','|---|---:|---:|---:|']
    for r in dist:
        md.append('| '+ ' | '.join([ZH.get(r['project_label_v0'],r['project_label_v0']),r['train'],r['val'],r['test']])+' |')
    md += ['', '测试集多数类（移动）占 132/217 = 60.83%；始终预测移动即可得到该准确率。因此同时报告 Macro-F1 和各类召回率。饮水仅2条、梳理仅3条，逐类结论不稳定。', '',
        '## 主对比表', '', '| '+' | '.join(headers)+' |','|'+'|'.join(['---']+['---:']*(len(headers)-1))+'|']
    md += ['| '+' | '.join(row)+' |' for row in table]
    if good:
        best=good[0];best_acc=max(good,key=lambda x:x['classification']['top1'])
        zero=[x['label'] for x in best['classification']['per_class'] if x['recall']==0]
        md += ['',f'本批次 Macro-F1 最高：{NAMES.get(best["model_id"],best["model_id"])}，{best["classification"]["macro_f1"]*100:.2f}%；Top-1 最高：{NAMES.get(best_acc["model_id"],best_acc["model_id"])}，{best_acc["classification"]["top1"]*100:.2f}%。',
               f'Macro-F1 最优模型的零召回类别：{", ".join(zero) if zero else "无"}。这只是当前 checkpoint 与当前测试划分的结果，不代表所有训练配方下的架构上限。']
    md += ['', '† FLOPs 使用 MMEngine JIT 统计，一个乘加计为一次运算；范围是输入对应的 backbone + neck + 分类头。带 * 的项目存在未支持算子，所报值是已支持算子的统计，不能当作精确完整 FLOPs。具体清单见各模型 `result.json`。', '',
        '## 测量口径', '',
        '- 每个模型使用现有 `best` 权重；同一模型有多次实验时，按历史验证 Top-1 选择。保留全部候选记录。历史训练配方/训练轮数不同，本表比较的是现有 checkpoint，而非严格同预算重训练。',
        '- 从训练目录保存的配置恢复模型结构，强制七类，并执行 `load_state_dict(strict=True)`；类别头不匹配直接失败。',
        '- 单中心裁剪；3D 模型一个时序 clip，2D 模型保留架构所需 segments，DenseSampleFrames 一个采样位置。各模型保留自身分辨率、帧数、采样间隔及归一化，不能把本表解释为完全相同输入算量。',
        '- batch=1，FP32，关闭 TF32，seed=42；每个模型独立进程，四个 CPU 线程、Decord 单线程。模型输入、原始/实际 pipeline 存于各模型目录。',
        '- 纯前向：输入已经解码、归一化并驻留 GPU；包含 backbone、neck、分类头，排除解码、传输及最终预测聚合。20次预热、100次计时；主表为同步主机计时，另保存 CUDA Event 计时。吞吐量按每视频采样视角计算，不是帧/秒。',
        '- 整体延迟：217条视频逐条解码、预处理、传输、模型预测、聚合以及结果拷回 CPU；包含文件读取，不包含模型加载和结果写盘。顺序访问会受到系统文件缓存影响，未清空操作系统缓存。',
        '- Top-1、Macro-F1、各类 Precision/Recall/F1 和混淆矩阵使用全部217条视频；P95 为重复测量的95分位数。Wilson区间仅表示测试样本量的不确定性，不包括重新训练的随机性。',
        '- 未测事件时间轴、异常漏检率或报警延迟；单标签片段分类指标不能替代这些系统指标。',
        '- TRN 模型内部关系采样含随机过程，使用固定随机种子，并在 FLOPs tracing 后重新设置种子。', '',
        '## 环境', '', '```json',json.dumps(env,ensure_ascii=False,indent=2),'```', '',
        '## 逐类别结果', '', '`per_class_all_models.csv` 与 Excel 的 per_class 工作表包含全部模型的各类 Precision/Recall/F1/support。', '',
        '![类别召回率](class_recall.png)', '', '![准确性与前向延迟](accuracy_latency.png)', '',
        '## 失败和未评测项目', '']
    for r in other: md.append(f'- `{r["model_id"]}`：{r["status"]}；{r.get("error","见execution.log")}')
    for r in manifest['exclusions']: md.append(f'- `{r["model_id"]}`：{r["reason"]}')
    md += ['', '不能用猫五分类微调权重、未微调的400类预训练权重或随机分类头补齐七类比较。', '',
        '## 文件说明', '',
        '- `comparison.xlsx`：总表、逐类别指标、训练元数据、数据分布、未评测项目。',
        '- `comparison.csv` / `comparison.json`：机器可读的主表/完整结果。',
        '- `comparison.html`：无需联网，可搜索和点击表头排序的报告。',
        '- `models/<model_id>/predictions.jsonl`：每条视频标签、预测、七类分数及耗时。',
        '- `models/<model_id>/confusion_matrix.json`、`per_class.csv`：混淆矩阵和逐类指标。',
        '- `models/<model_id>/protocol.json`、`resolved_config.py`、`forward_timings.json`：实际配置和原始计时。',
        '- `models/<model_id>/execution.log`、`result.json`：日志、加载验证、复杂度警告及完整指标。',
        '- `manifest.json`、`environment.json`、`dataset/`：权重选择、环境、数据来源和固定划分证据。',
        '- `scripts/`：本批次脚本快照；本地仓库 `scripts/` 下亦有源脚本。远程项目原有跟踪源码未修改。',
        '- `report_audit.json`：从逐视频预测独立重算样本数、Top-1 和混淆矩阵的校验。', '',
        '## 复现', '', '在远程项目根目录、同一环境中运行（新目录会完整重测，已有目录会跳过成功项）：', '',
        '```bash',
        '~/miniconda3/envs/pet/bin/python -B scripts/compare_mammal_models.py \\',
        '  --repo /home/wyy/pet-action-recognition \\',
        '  --out /home/wyy/pet-action-recognition/results/model_comparison/NEW_BATCH \\',
        '  --device cuda:0 --warmup 20 --repetitions 100',
        '```', '', '远程使用本批次源快照时，将上述脚本路径替换为本结果目录下的 `scripts/compare_mammal_models_v2.py`。v2增加了TPN懒创建层和项目自定义损失的加载兼容，不更改预测协议；首轮环境记录另存 `environment_full_v1.json`。']
    (p/'REPORT.md').write_text('\n'.join(md)+'\n',encoding='utf-8')

    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import numpy as np
    plt.rcParams['font.family']='DejaVu Sans'
    if good:
        fig,ax=plt.subplots(figsize=(10, max(5,len(good)*.38)))
        arr=np.array([[x['recall'] for x in r['classification']['per_class']] for r in good])
        im=ax.imshow(arr,vmin=0,vmax=1,cmap='YlGnBu',aspect='auto')
        ax.set_xticks(range(7),['Move\n(n=132)','Jump\n(n=14)','Eat\n(n=11)','Drink\n(n=2)','Groom\n(n=3)','Rest\n(n=32)','Social\n(n=23)'])
        ax.set_yticks(range(len(good)),[NAMES.get(r['model_id'],r['model_id']) for r in good])
        for i in range(len(good)):
            for j in range(7): ax.text(j,i,f'{arr[i,j]*100:.0f}%',ha='center',va='center',fontsize=8,color='white' if arr[i,j]>.6 else 'black')
        ax.set_title('Per-class recall | AnimalKingdom-derived mammal test (217 clips)')
        fig.colorbar(im,ax=ax,label='Recall');fig.tight_layout();fig.savefig(p/'class_recall.png',dpi=180);plt.close(fig)
        fig,ax=plt.subplots(figsize=(11,7))
        for i,r in enumerate(good):
            x=r['forward_host']['mean_ms'];y=r['classification']['macro_f1']*100
            ax.scatter(x,y,s=45)
            ax.annotate(NAMES.get(r['model_id'],r['model_id']).replace(' [ID: swin-tiny]',''),(x,y),xytext=(4,5 if i%2 else -10),textcoords='offset points',fontsize=8)
        ax.set_xlabel('Synchronized forward latency (ms, batch 1, FP32)');ax.set_ylabel('Macro-F1 (%)')
        ax.set_title('Existing fine-tuned checkpoints | single-view inference');ax.grid(alpha=.2)
        fig.tight_layout();fig.savefig(p/'accuracy_latency.png',dpi=180);plt.close(fig)
    # HTML contains the full report as escaped text plus the interactive table.
    def embedded(name):
        path=p/name
        return 'data:image/png;base64,'+base64.b64encode(path.read_bytes()).decode() if path.exists() else ''
    tr=''.join('<tr>'+''.join('<td>'+html.escape(v)+'</td>' for v in row)+'</tr>' for row in table)
    head=''.join('<th onclick="sortTable('+str(i)+')">'+html.escape(v)+'</th>' for i,v in enumerate(headers))
    page='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>AnimalKingdom 哺乳动物子集模型对比</title><style>
body{font:15px/1.6 system-ui,sans-serif;margin:32px;color:#172b3a;background:#f5f7fa}h1{font-size:27px}table{border-collapse:collapse;width:100%;background:white}th,td{padding:10px;border-bottom:1px solid #ddd;text-align:right}th{cursor:pointer;background:#dfebf3;position:sticky;top:0}td:first-child,th:first-child{text-align:left}input{padding:10px;margin:14px 0;width:320px}img{max-width:100%;background:white}pre{white-space:pre-wrap;background:white;padding:24px;border:1px solid #ddd}section{overflow:auto}small{color:#536777}</style>
<h1>AnimalKingdom 哺乳动物子集模型对比</h1><p>七类 · 217 条测试视频 · 现有微调权重 · 单视角 FP32 / batch=1</p>
<small>点击表头排序。GFLOPs 带 * 表示存在未统计算子；架构输入规格不同，详见报告。未评测原始 AnimalKingdom 全类别任务。</small>
<br><input id="q" placeholder="搜索模型" oninput="filterRows()"><section><table id="t"><thead><tr>'''+head+'''</tr></thead><tbody>'''+tr+'''</tbody></table></section>
<h2>各类别召回率</h2><img src="class_recall.png"><h2>效果与延迟</h2><img src="accuracy_latency.png">
<h2>完整说明与复现信息</h2><pre>'''+html.escape('\n'.join(md))+'''</pre><script>
let direction={};function sortTable(i){let b=document.querySelector('#t tbody');let rs=[...b.rows];direction[i]=-(direction[i]||-1);rs.sort((a,b)=>{let x=a.cells[i].textContent,y=b.cells[i].textContent;let n=parseFloat(x),m=parseFloat(y);return direction[i]*(Number.isNaN(n)||Number.isNaN(m)?x.localeCompare(y):n-m)});rs.forEach(r=>b.appendChild(r));}
function filterRows(){let q=document.querySelector('#q').value.toLowerCase();document.querySelectorAll('#t tbody tr').forEach(r=>r.hidden=!r.textContent.toLowerCase().includes(q));}</script></html>'''
    page=page.replace('src="class_recall.png"','src="'+embedded('class_recall.png')+'"').replace('src="accuracy_latency.png"','src="'+embedded('accuracy_latency.png')+'"')
    (p/'comparison.html').write_text(page,encoding='utf-8')
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    wb=Workbook();wb.remove(wb.active)
    def sheet(name, rows):
        ws=wb.create_sheet(name)
        if not rows:return
        keys=list(dict.fromkeys(k for r in rows for k in r))
        ws.append(keys)
        for r in rows:
            ws.append([json.dumps(r.get(k),ensure_ascii=False) if isinstance(r.get(k),(dict,list)) else r.get(k) for k in keys])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='284E70')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(55,max(14,max(len(str(c.value or '')) for c in list(col)[:30])+2))
    flat=list(csv.DictReader((p/'comparison.csv').open(encoding='utf-8-sig')))
    for row in flat:
        for key,value in list(row.items()):
            if key not in ['model','status','input_shape','unsupported_flop_ops','error','run_id'] and value:
                try:row[key]=float(value)
                except ValueError:pass
    sheet('comparison',flat);sheet('per_class',all_class);sheet('dataset',dist)
    sheet('training',[dict(model=r['model_id'],**r.get('historical_training',{})) for r in good])
    sheet('exclusions',manifest['exclusions']+[dict(model_id=r['model_id'],reason=r.get('error',r['status'])) for r in other])
    sheet('protocol',[dict(item='说明',value='Top-1、Macro-F1、Recall 为0至1小数；数据是AnimalKingdom映射七类哺乳动物子集，非原始官方全类别基准。详见 REPORT.md。')])
    wb.save(p/'comparison.xlsx')
    print(json.dumps(dict(completed=len(good),other=len(other),best_macro_f1=good[0]['model_id'] if good else None),ensure_ascii=False))

if __name__=='__main__':main()
